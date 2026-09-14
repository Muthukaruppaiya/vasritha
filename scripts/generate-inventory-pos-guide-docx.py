"""Generate docs/INVENTORY_AND_POS_GUIDE.docx from the user guide content."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "INVENTORY_AND_POS_GUIDE.docx"

BRAND = RGBColor(0x7A, 0x4F, 0x2B)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
        set_cell_shading(hdr[i], "F8F1EA")
    for ri, row in enumerate(rows):
        for ci, text in enumerate(row):
            table.rows[ri + 1].cells[ci].text = text
    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def main() -> None:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("Inventory, Inward Stock & Store POS", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = BRAND

    sub = doc.add_paragraph("User guide for shop staff and managers")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = MUTED
    sub.runs[0].italic = True

    doc.add_paragraph(
        "Sukadhaa / Vasritha admin (operations panel)"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("Sign in: ").bold = True
    p.add_run("Open your admin URL (e.g. /admin/login) and use your staff email and password.")

    doc.add_heading("Before you start", level=1)
    doc.add_paragraph(
        "Complete these once, or when adding a new product:"
    )
    add_table(
        doc,
        ["Step", "Where", "What to do"],
        [
            [
                "1",
                "Catalogue → Categories",
                "Create the category (and subcategory if needed), e.g. Sarees → Cotton sarees",
            ],
            [
                "2",
                "Catalogue → Product Master",
                "Create the product: name, price, category, status Active",
            ],
            [
                "3",
                "System → Shops",
                "Ensure at least one shop is active (required for store billing)",
            ],
        ],
    )
    note = doc.add_paragraph()
    note.add_run("Remember: ").bold = True
    note.add_run(
        "Product details are edited in Product Master. Stock quantities change only in "
        "Inventory (receive, adjust) or when a sale happens in POS / online."
    )

    doc.add_heading("End-to-end flow (quick view)", level=1)
    flow = doc.add_paragraph()
    flow.add_run(
        "Categories  →  Product Master  →  Inventory (receive stock)  →  "
        "Print barcodes  →  Store POS (bill)"
    ).bold = True
    add_numbered(
        doc,
        [
            "Create the product in Product Master.",
            "Receive stock when goods arrive (GRN / inward).",
            "Print barcode stickers for each unique piece (from Product Master).",
            "Bill the customer in Store POS (scan barcode or search by name/SKU).",
        ],
    )

    doc.add_heading("1. Inventory flow (view & check stock)", level=1)
    doc.add_paragraph().add_run("Menu: Catalogue → Inventory (/admin/inventory)").bold = True
    doc.add_paragraph(
        "Use this page to see how much stock you have and what moved recently."
    )

    doc.add_heading("What you see", level=2)
    add_bullets(
        doc,
        [
            "Summary cards — total SKUs, total pieces on hand, in stock / low / out of stock",
            "Current stock table — each product variant with On hand count and level badge",
            "Stock ledger — recent movements (inward, sale, adjustment, etc.)",
        ],
    )

    doc.add_heading("Stock level badges", level=2)
    add_table(
        doc,
        ["Badge", "Meaning"],
        [
            ["In stock", "Enough quantity available"],
            ["Low stock", "On hand is 10 or less (system default)"],
            ["Out of stock", "Zero pieces available"],
        ],
    )

    doc.add_heading("Filters", level=2)
    add_bullets(
        doc,
        [
            "Search — product name, SKU, HSN, category",
            "Stock level — All / In stock / Low stock / Out of stock",
            "Category / Subcategory — narrow the list",
        ],
    )

    doc.add_heading("Row actions", level=2)
    add_table(
        doc,
        ["Button", "Use when"],
        [
            ["Receive", "New stock arrived — opens the inward (GRN) form for that variant"],
            [
                "Adjust",
                "Correction only (count mistake, damage, return to shelf) — not for new purchases",
            ],
            ["Edit", "Opens Product Master to change product details (not quantity)"],
        ],
    )

    doc.add_heading("From Product Master", level=2)
    doc.add_paragraph(
        "On a product row, click Stock — you are taken to Inventory filtered for that product."
    )

    doc.add_heading("2. Inward flow (receive stock / GRN)", level=1)
    doc.add_paragraph().add_run("When to use: ").bold = True
    doc.add_paragraph(
        "Goods arrive from a supplier, warehouse, or production. This adds stock and "
        "creates unique barcode tags for each piece."
    )
    doc.add_paragraph().add_run(
        "Where: Catalogue → Inventory → Receive stock (top button) or Receive on a row."
    ).bold = True

    doc.add_heading("Step-by-step", level=2)
    add_numbered(
        doc,
        [
            "Open Receive stock (GRN).",
            "Optional: Fill Supplier, Bill / invoice no., and Note.",
            "Under Lines: choose SKU / variant, enter Quantity, use + Add line for multiple SKUs.",
            "Click Post inward.",
            "Wait for success — stock table and ledger update with Inward / purchase.",
        ],
    )

    doc.add_heading("After inward", level=2)
    add_numbered(
        doc,
        [
            "Go to Product Master → open the product → print barcode stickers for new pieces.",
            "Stick labels on each physical item before display.",
            "Confirm On hand on Inventory matches what you received.",
        ],
    )

    doc.add_heading("Opening stock (new product only)", level=2)
    doc.add_paragraph(
        "When creating a product in Product Master, you may set Opening stock (unique pieces) "
        "once for the first batch. After save, do not change quantity on the product form — "
        "use Inventory → Receive stock for all future inward."
    )

    doc.add_heading("Adjust stock (not inward)", level=2)
    doc.add_paragraph("Use Adjust stock only for corrections:")
    add_table(
        doc,
        ["Type", "Use for"],
        [
            ["Opening stock (+)", "Rare one-time correction"],
            ["Customer / supplier return (+)", "Item came back to sellable stock"],
            ["Manual correction (+/−)", "Stock count, damage write-off, etc."],
        ],
    )
    doc.add_paragraph(
        "Always add a Note / reason (required for audits). "
        "Prefer Receive (GRN) for all normal purchases from suppliers."
    )

    doc.add_heading("3. Store POS billing flow", level=1)
    doc.add_paragraph().add_run("Menu: Sales → Store POS (/admin/billing)").bold = True
    doc.add_paragraph("Use this to bill walk-in customers at the shop counter.")

    doc.add_heading("Before billing", level=2)
    add_bullets(
        doc,
        [
            "Shop selected in the Shop dropdown (top of page).",
            "Products are Active and have stock > 0.",
            "Barcode stickers are printed and on the items (recommended).",
        ],
    )

    doc.add_heading("Step-by-step sale", level=2)

    doc.add_heading("A. Add items to cart", level=3)
    add_numbered(
        doc,
        [
            "Click in Scan barcode or search name / SKU.",
            "Scan a piece barcode or type product name / family SKU and press Enter or Add.",
            "If several products match, pick the correct one from the list.",
            "Repeat for each item.",
        ],
    )
    add_bullets(
        doc,
        [
            "Scanning a unique piece barcode adds exactly one piece (qty locked to 1).",
            "Searching by family SKU may allow quantity + / − up to available stock.",
        ],
    )

    doc.add_heading("B. Customer details (required)", level=3)
    add_table(
        doc,
        ["Field", "Required?"],
        [
            ["Name", "Yes"],
            ["Mobile", "Yes (10-digit Indian mobile)"],
            ["Email", "No"],
        ],
    )

    doc.add_heading("C. Discount (optional)", level=3)
    add_bullets(
        doc,
        [
            "Choose % Off or ₹ Off, then enter the discount amount.",
            "Check Payable total before payment.",
        ],
    )

    doc.add_heading("D. Payment", level=3)
    add_table(
        doc,
        ["Method", "Action"],
        [
            ["Cash", "Click Collect ₹… — bill is paid immediately"],
            [
                "Razorpay",
                "Click Pay ₹… with Razorpay — customer completes UPI/card on Razorpay screen",
            ],
        ],
    )
    doc.add_paragraph("On success, a bill preview opens.")

    doc.add_heading("E. Print & finish", level=3)
    add_numbered(
        doc,
        [
            "Click Print bill (5″ thermal-style shop bill).",
            "Click Close to start the next sale, or New sale in the header to clear the cart.",
        ],
    )
    doc.add_paragraph(
        "Stock is reduced automatically when payment is paid. "
        "Each sold unique piece is marked so it cannot be sold again."
    )

    doc.add_heading("Reprint a past bill", level=2)
    doc.add_paragraph().add_run("Menu: Invoice → Store Invoice (/admin/invoices/store)").bold = True
    add_bullets(
        doc,
        [
            "Search by invoice number, customer name, phone, or email.",
            "Open a paid POS bill and print again.",
        ],
    )

    doc.add_heading("If Razorpay is cancelled", level=2)
    doc.add_paragraph(
        "The order stays pending. You can retry Razorpay or switch to Cash for the same cart."
    )

    doc.add_heading("Who can do what (summary)", level=1)
    add_table(
        doc,
        ["Task", "Typical roles"],
        [
            ["View inventory", "Inventory staff, manager, owner"],
            ["Receive stock (GRN)", "Inventory staff, manager, owner"],
            ["Adjust stock (manual)", "Manager, owner (needs approval permission)"],
            ["Store POS billing", "Billing staff, manager, owner"],
            ["Manage products", "Manager, owner"],
        ],
    )
    doc.add_paragraph(
        "If a menu or button is missing, ask your admin to check System → Users and role permissions."
    )

    doc.add_heading("Common issues", level=1)
    add_table(
        doc,
        ["Problem", "What to check"],
        [
            ["Product not found in POS", "Product status Active? Stock > 0? Inward posted?"],
            [
                "Barcode scan not working",
                "Scan the piece tag (e.g. VAS-…-0001), not only the family code",
            ],
            ["No active shop", "System → Shops — add or activate a shop"],
            ["Stock wrong after sale", "Inventory → Stock ledger for Sale out lines"],
            ["Need more stock", "Inventory → Receive stock, not Product Master edit"],
        ],
    )

    doc.add_heading("Daily checklist (shop floor)", level=1)
    checklist = [
        "Receive inward (GRN) when new cartons arrive",
        "Print and attach barcodes before display",
        "Bill in Store POS with customer name + mobile",
        "Glance at Low stock filter on Inventory before closing",
    ]
    for item in checklist:
        doc.add_paragraph(f"☐  {item}")

    doc.add_paragraph()
    foot = doc.add_paragraph(
        "Last updated for the current Vasritha release. "
        "For technical setup, see CURRENT_RELEASE.md and LOCAL_POSTGRES.md."
    )
    foot.runs[0].italic = True
    foot.runs[0].font.size = Pt(9)
    foot.runs[0].font.color.rgb = MUTED

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
